"""
Module 7 - Report Export (Excel with CSV fallback)
==================================================

Exports three sheets: Summary (with the 75 % rule), Attendance and Students.
If openpyxl is not installed the export does not fail: it falls back to
plain CSV files and says so (graceful degradation).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from .analytics import THRESHOLD

STATUS_FILL = {"PRESENT": "DCFCE7", "LATE": "FEF3C7", "ABSENT": "FEE2E2",
               "EXCUSED": "EDE9FE", "ON_DUTY": "CFFAFE", "NOT_MARKED": "F1F5F9"}
SUMMARY_HEADER = ["Reg No", "Name", "Held", "Present", "Late", "Absent", "Excused",
                  "On Duty", "Attendance %", "Status", "Classes needed for 75%"]


@dataclass
class ExportResult:
    files: list
    used_excel: bool
    message: str


def _summary_rows(summaries):
    for s in summaries:
        pct = s.percentage
        yield [s.regno, s.name, s.held, s.present, s.late, s.absent, s.excused, s.on_duty,
               "" if pct is None else round(pct, 1), s.status, s.classes_needed]


def export_report(path, students, attendance_rows, summaries, prefer_excel=True) -> ExportResult:
    path = Path(path)
    if prefer_excel:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            prefer_excel = False
            reason = "openpyxl is not installed, so the report was saved as CSV files instead."
        else:
            reason = ""
    else:
        reason = "Excel export disabled, report saved as CSV files."

    if prefer_excel:
        _write_excel(path.with_suffix(".xlsx"), students, attendance_rows, summaries)
        return ExportResult([path.with_suffix(".xlsx")], True,
                            f"Report exported to {path.with_suffix('.xlsx').name}")

    files = []
    stem = path.with_suffix("")
    tables = {
        "summary": (SUMMARY_HEADER, list(_summary_rows(summaries))),
        "attendance": (["SessionID", "Date", "RegNo", "Name", "Status", "MarkedAt"],
                       [[r.get(k, "") for k in ("SessionID", "Date", "RegNo", "Name", "Status",
                                                "MarkedAt")] for r in attendance_rows]),
        "students": (["RegNo", "Name", "DateAdded"],
                     [[s.regno, s.name, s.added] for s in students.values()]),
    }
    for name, (header, rows) in tables.items():
        out = Path(f"{stem}_{name}.csv")
        with open(out, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(header)
            writer.writerows(rows)
        files.append(out)
    return ExportResult(files, False, f"{reason} Saved {', '.join(f.name for f in files)}")


def _write_excel(path: Path, students, attendance_rows, summaries) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", start_color="1E3A8A", end_color="1E3A8A")
    header_font = Font(bold=True, color="FFFFFF")
    low_fill = PatternFill("solid", start_color="FEE2E2", end_color="FEE2E2")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    sheet(ws, SUMMARY_HEADER, list(_summary_rows(summaries)),
          [13, 22, 7, 9, 7, 8, 9, 9, 13, 13, 22])
    for row in ws.iter_rows(min_row=2):
        pct = row[8].value
        if pct != "" and pct is not None and float(pct) < THRESHOLD:
            for cell in row:
                cell.fill = low_fill

    ws2 = wb.create_sheet("Attendance")
    keys = ["SessionID", "Date", "RegNo", "Name", "Status", "MarkedAt"]
    sheet(ws2, keys, [[r.get(k, "") for k in keys] for r in attendance_rows],
          [18, 12, 13, 22, 13, 11])
    for row in ws2.iter_rows(min_row=2):
        color = STATUS_FILL.get(row[4].value)
        if color:
            row[4].fill = PatternFill("solid", start_color=color, end_color=color)

    ws3 = wb.create_sheet("Students")
    sheet(ws3, ["RegNo", "Name", "DateAdded"],
          [[s.regno, s.name, s.added] for s in students.values()], [13, 24, 18])
    wb.save(path)
